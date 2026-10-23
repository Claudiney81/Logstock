import sys
import unittest
import io
from unittest.mock import patch
from openpyxl import load_workbook
from pathlib import Path
from datetime import datetime, date, timedelta, timezone

sys.path.insert(0, str(Path(__file__).parent))
from test_estoque_minimos import EstoqueMinimosTests
from app.models import NotaFiscalEntrada, Estoque, NotaFiscalEdicao
from app.extensions import db
from app.utils.datas import formatar_horario, hoje_local, limites_dia_utc
from app.utils.edicao_nota import assinatura


class NotaDatasTests(unittest.TestCase):
    setUp = EstoqueMinimosTests.setUp
    tearDown = EstoqueMinimosTests.tearDown

    def dados(self):
        return {'numero_nf': 'DATA-TESTE', 'fornecedor': 'Fornecedor Teste',
                'tipo_estoque': 'empresa', 'tipo_servico_id': '1',
                'categoria_entrada': 'MATERIAL', 'codigo[]': 'TESTE01',
                'descricao[]': 'Cabo teste', 'quantidade[]': '2', 'valor[]': '12,50',
                'data_recebimento': (hoje_local() - timedelta(days=3)).isoformat()}

    def test_receipt_separate_from_registration_and_stock(self):
        antes = Estoque.query.filter_by(item_id=1).one().quantidade
        data = self.dados()
        self.client.post('/nota/nova', data=data)
        nota = NotaFiscalEntrada.query.filter_by(numero_nf='DATA-TESTE').one()
        self.assertEqual(nota.data_recebimento.isoformat(), data['data_recebimento'])
        self.assertLess(abs((datetime.utcnow() - nota.data_hora).total_seconds()), 30)
        self.assertEqual(Estoque.query.filter_by(item_id=1).one().quantidade, antes + 2)
        registro = nota.data_hora
        self.client.post(f'/nota/{nota.id}/editar', data={
            'versao': assinatura(nota), 'numero_nf': nota.numero_nf,
            'fornecedor': nota.fornecedor, 'observacao': nota.observacao,
            'data_recebimento': (hoje_local() - timedelta(days=4)).isoformat(),
            'motivo': 'Conferência da entrega',
        })
        db.session.expire_all()
        self.assertEqual(nota.data_hora, registro)
        self.assertEqual(Estoque.query.filter_by(item_id=1).one().quantidade, antes + 2)
        self.assertEqual(NotaFiscalEdicao.query.count(), 1)
        for path in ['/nota/historico', f'/nota/{nota.id}', '/nota/nova']:
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200)
            self.assertIn('recebimento', response.get_data(as_text=True).lower())

    def test_future_and_invalid_receipt_rejected_before_stock_write(self):
        antes = Estoque.query.filter_by(item_id=1).one().quantidade
        for valor in ['2026-99-99', (hoje_local() + timedelta(days=1)).isoformat()]:
            data = self.dados(); data['data_recebimento'] = valor
            self.client.post('/nota/nova', data=data)
            self.assertEqual(NotaFiscalEntrada.query.count(), 0)
            self.assertEqual(Estoque.query.filter_by(item_id=1).one().quantidade, antes)

    def test_local_timezone_and_midnight_boundaries(self):
        self.assertEqual(formatar_horario(datetime(2026, 10, 23, 13, 15)), '23/10/2026 10:15')
        self.assertEqual(formatar_horario(datetime(2026, 10, 23, 1)), '22/10/2026 22:00')
        self.assertEqual(formatar_horario(datetime(2026, 10, 23, 13, 15, tzinfo=timezone.utc)), '23/10/2026 10:15')
        self.assertEqual(limites_dia_utc(date(2026, 10, 23)),
                         (datetime(2026, 10, 23, 3), datetime(2026, 10, 24, 3)))

    def test_receipt_reflected_in_excel_pdf_and_value_history(self):
        dados = self.dados()
        self.client.post('/nota/nova', data=dados)
        nota = NotaFiscalEntrada.query.filter_by(numero_nf='DATA-TESTE').one()
        esperado = nota.data_recebimento.strftime('%d/%m/%Y')
        response = self.client.get(f'/nota/{nota.id}/excel')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.data), data_only=True).active
        linhas = list(sheet.values)
        self.assertTrue(any('Recebimento' in row and esperado in row for row in linhas))
        # Validate the actual HTML passed to the existing PDF renderer.
        with patch('app.routes.nota_fiscal.pdfkit.from_string', return_value=b'%PDF-test') as renderer:
            response = self.client.get(f'/nota/{nota.id}/pdf')
            self.assertEqual(response.status_code, 200)
            html = renderer.call_args.args[0]
            self.assertIn(esperado, html)
            self.assertIn(nota.fornecedor, html)
            self.assertIn(formatar_horario(nota.data_hora), html)
        response = self.client.get('/nota/historico-valores')
        self.assertEqual(response.status_code, 200)
        self.assertIn(esperado, response.get_data(as_text=True))
        response = self.client.get('/nota/historico-valores/excel')
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data), data_only=True)
        self.assertTrue(any(esperado == cell.value for sheet in workbook for row in sheet for cell in row))
